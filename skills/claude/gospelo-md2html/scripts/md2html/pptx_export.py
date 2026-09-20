"""PPTX export: one image slide per rendered page, plus invisible link hotspots
and an optional transparent selectable-text layer.

The slide is a screenshot of the page's <section class="page"> taken by the
same Chromium that verified the layout, so the deck shows exactly what the
HTML shows. Links and text lines are read from the DOM (fractions of the page
box), so no PDF round trip is needed. Slide size equals the page format.
"""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from .measure import Browser

MM_TO_EMU = 36000
PT_PER_MM = 72.0 / 25.4

_PAGES_JS = r"""
(withText) => Array.from(document.querySelectorAll('section.page')).map(sec => {
  const R = sec.getBoundingClientRect();
  const rel = (r) => [(r.left - R.left) / R.width, (r.top - R.top) / R.height, r.width / R.width, r.height / R.height];
  const links = [];
  for (const a of sec.querySelectorAll('a[href]')) {
    for (const r of a.getClientRects()) {
      if (r.width > 0 && r.height > 0) { links.push({ href: a.href, rect: rel(r) }); }
    }
  }
  const lines = [];
  if (withText) {
    const walker = document.createTreeWalker(sec, NodeFilter.SHOW_TEXT);
    let node;
    while ((node = walker.nextNode())) {
      const text = node.nodeValue;
      if (!text.trim()) { continue; }
      const p = node.parentElement;
      if (!p || p.closest('svg, script, style')) { continue; }
      const cs = getComputedStyle(p);
      if (cs.visibility === 'hidden' || cs.display === 'none') { continue; }
      const size = parseFloat(cs.fontSize);
      let cur = null, lastTop = null;
      for (let i = 0; i < text.length; i++) {
        const rg = document.createRange();
        rg.setStart(node, i); rg.setEnd(node, i + 1);
        const rects = rg.getClientRects();
        if (!rects.length) { continue; }
        const rc = rects[0];
        if (rc.width === 0 && text[i] === '\n') { continue; }
        if (cur === null || Math.abs(rc.top - lastTop) > 1) {
          if (cur !== null) { lines.push(cur); }
          cur = { text: '', left: rc.left, top: rc.top, right: rc.right, bottom: rc.bottom, size };
          lastTop = rc.top;
        }
        cur.text += text[i];
        cur.left = Math.min(cur.left, rc.left); cur.right = Math.max(cur.right, rc.right);
        cur.top = Math.min(cur.top, rc.top); cur.bottom = Math.max(cur.bottom, rc.bottom);
      }
      if (cur !== null) { lines.push(cur); }
    }
  }
  return {
    id: sec.dataset.pageId,
    links,
    lines: lines.filter(l => l.text.trim()).map(l => ({
      text: l.text, size: l.size / R.height,
      rect: rel({ left: l.left, top: l.top, width: l.right - l.left, height: l.bottom - l.top })
    }))
  };
})
"""


@dataclass
class SlideData:
    id: str
    image: bytes
    image_format: str                       # "jpeg" | "png"
    links: list[dict[str, Any]] = field(default_factory=list)   # {href, rect: [x, y, w, h] fractions}
    lines: list[dict[str, Any]] = field(default_factory=list)   # {text, size (fraction of page height), rect}


@dataclass(frozen=True)
class PptxOptions:
    dpi: int = 200                 # screenshot resolution (CSS px are 96 dpi)
    image_format: str = "jpeg"     # "jpeg" | "png"
    jpeg_quality: int = 92
    text_layer: bool = False       # transparent selectable text over the image
    link_hotspots: bool = True     # invisible clickable rectangles over <a href>


def capture_pages(browser: Browser, html_path: Path, opts: PptxOptions) -> list[SlideData]:
    """Screenshot every rendered page and read its links (and text lines) from the DOM."""
    page = browser.open(html_path, scale=opts.dpi / 96.0)
    try:
        meta = page.evaluate(_PAGES_JS, opts.text_layer)
        sections = page.query_selector_all("section.page")
        if len(sections) != len(meta):
            raise RuntimeError("page count mismatch while capturing slides")
        out = []
        for sec, m in zip(sections, meta):
            if opts.image_format == "png":
                img = sec.screenshot(type="png")
            else:
                img = sec.screenshot(type="jpeg", quality=opts.jpeg_quality)
            out.append(SlideData(m["id"], img, opts.image_format, m["links"], m["lines"]))
        return out
    finally:
        page.close()


def assemble_pptx(slides: list[SlideData], width_mm: float, height_mm: float, pptx_path: Path,
                  opts: PptxOptions) -> dict[str, int]:
    """Write the deck. Returns counts for the report."""
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.oxml.ns import qn
    from pptx.util import Emu, Pt

    slide_w = int(round(width_mm * MM_TO_EMU))
    slide_h = int(round(height_mm * MM_TO_EMU))
    page_h_pt = height_mm * PT_PER_MM

    prs = Presentation()
    prs.slide_width = slide_w
    prs.slide_height = slide_h
    blank = prs.slide_layouts[6]
    n_links = n_lines = 0

    def box(rect):
        x, y, w, h = rect
        return (Emu(int(x * slide_w)), Emu(int(y * slide_h)),
                Emu(max(1, int(w * slide_w))), Emu(max(1, int(h * slide_h))))

    def transparent(fill_or_run_xml):
        srgb = fill_or_run_xml.find(qn("a:solidFill")).find(qn("a:srgbClr"))
        srgb.append(srgb.makeelement(qn("a:alpha"), {"val": "0"}))

    for s in slides:
        slide = prs.slides.add_slide(blank)
        slide.shapes.add_picture(io.BytesIO(s.image), Emu(0), Emu(0), Emu(slide_w), Emu(slide_h))
        if opts.text_layer:
            for ln in s.lines:
                tb = slide.shapes.add_textbox(*box(ln["rect"]))
                tf = tb.text_frame
                tf.word_wrap = False
                tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
                run = tf.paragraphs[0].add_run()
                run.text = ln["text"]
                run.font.size = Pt(max(1.0, ln["size"] * page_h_pt))
                run.font.color.rgb = RGBColor(0, 0, 0)
                rPr = run._r.get_or_add_rPr()
                rPr.set("noProof", "1")   # no spell-check squiggle over invisible text
                transparent(rPr)
                n_lines += 1
        if opts.link_hotspots:
            for lk in s.links:
                shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, *box(lk["rect"]))
                shp.fill.solid()
                shp.fill.fore_color.rgb = RGBColor(0, 0, 0)
                transparent(shp.fill._xPr)   # clickable but invisible
                shp.line.fill.background()
                shp.shadow.inherit = False
                shp.click_action.hyperlink.address = lk["href"]
                n_links += 1
    pptx_path.parent.mkdir(parents=True, exist_ok=True)
    prs.save(str(pptx_path))
    return {"slides": len(slides), "links": n_links, "textLines": n_lines}


def export_pptx(browser: Browser, html_path: Path, pptx_path: Path, width_mm: float, height_mm: float,
                opts: PptxOptions) -> dict[str, int]:
    slides = capture_pages(browser, html_path, opts)
    return assemble_pptx(slides, width_mm, height_mm, pptx_path, opts)
