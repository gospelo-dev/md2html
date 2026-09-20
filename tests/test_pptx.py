"""PPTX assembly (no browser): slide size, one picture per page, link hotspots, transparent text layer."""

import io
import sys
from pathlib import Path

import pytest

sys.path.insert(0, str(Path(__file__).resolve().parent.parent / "skills" / "claude" / "gospelo-md2html" / "scripts"))

pptx = pytest.importorskip("pptx")
PIL = pytest.importorskip("PIL")

from md2html.pptx_export import MM_TO_EMU, PptxOptions, SlideData, assemble_pptx  # noqa: E402


def png_bytes(w=64, h=36) -> bytes:
    from PIL import Image
    buf = io.BytesIO()
    Image.new("RGB", (w, h), (240, 240, 240)).save(buf, format="PNG")
    return buf.getvalue()


def slides():
    return [
        SlideData("p01", png_bytes(), "png",
                  links=[{"href": "https://example.com/", "rect": [0.1, 0.2, 0.3, 0.05]}],
                  lines=[{"text": "Hello", "size": 0.03, "rect": [0.1, 0.1, 0.2, 0.04]}]),
        SlideData("p02", png_bytes(), "png"),
    ]


def test_deck_has_one_picture_per_page_and_the_page_size(tmp_path):
    out = tmp_path / "deck.pptx"
    counts = assemble_pptx(slides(), 338.6667, 190.5, out, PptxOptions(link_hotspots=True, text_layer=False))
    assert counts == {"slides": 2, "links": 1, "textLines": 0}
    prs = pptx.Presentation(str(out))
    assert prs.slide_width == round(338.6667 * MM_TO_EMU) and prs.slide_height == round(190.5 * MM_TO_EMU)
    assert len(prs.slides) == 2
    s1 = prs.slides[0]
    pics = [sh for sh in s1.shapes if sh.shape_type == 13]  # PICTURE
    assert len(pics) == 1 and pics[0].width == prs.slide_width
    hot = [sh for sh in s1.shapes if sh.click_action.hyperlink.address]
    assert len(hot) == 1 and hot[0].click_action.hyperlink.address == "https://example.com/"
    assert hot[0].left == int(0.1 * prs.slide_width)


def test_text_layer_is_transparent_and_not_proofed(tmp_path):
    out = tmp_path / "deck.pptx"
    counts = assemble_pptx(slides(), 210, 297, out, PptxOptions(text_layer=True, link_hotspots=False))
    assert counts["textLines"] == 1 and counts["links"] == 0
    prs = pptx.Presentation(str(out))
    boxes = [sh for sh in prs.slides[0].shapes if sh.has_text_frame and sh.text_frame.text == "Hello"]
    assert len(boxes) == 1
    run = boxes[0].text_frame.paragraphs[0].runs[0]
    rpr = run._r.get_or_add_rPr()
    assert rpr.get("noProof") == "1"
    from pptx.oxml.ns import qn
    assert rpr.find(qn("a:solidFill")).find(qn("a:srgbClr")).find(qn("a:alpha")).get("val") == "0"
    # 0.03 of a 297mm page is about 25pt
    assert abs(run.font.size.pt - 0.03 * 297 * 72 / 25.4) < 0.5
